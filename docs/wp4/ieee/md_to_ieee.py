"""Authored report markdown -> two-column IEEE-style .docx.
Custom markers: %author, %affil, %abstract, %keywords, %table key | caption | c|f, %eq key | text, %references,
figures ![fig:key|caption|c|f|path](path), cross references {sec:k} {secn:k} {fig:k} {tab:k} {app:X} {cite:N} {eq:k}."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1])
OUT = sys.argv[2]
ROOT = Path(sys.argv[3])  # folder the figure paths are relative to

PAGE_W, PAGE_H = 21.0, 29.7
MARGIN_LR = 1.6
GAP = 0.6
FULL_W = PAGE_W - 2 * MARGIN_LR          # 17.8 cm
COL_W = (FULL_W - GAP) / 2               # 8.6 cm
FONT = "Times New Roman"
MONO = "Courier New"
ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV",
         "XVI", "XVII", "XVIII", "XIX", "XX", "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI", "XXVII", "XXVIII",
         "XXIX", "XXX", "XXXI", "XXXII", "XXXIII", "XXXIV", "XXXV"]


# ---------------------------------------------------------------------------------------------- parsing
def parse(text):
    lines = text.split("\n")
    blocks, i, n = [], 0, len(lines)
    in_refs = False
    while i < n:
        raw = lines[i]
        st = raw.strip()
        if not st:
            i += 1
            continue
        if st.startswith("# "):
            blocks.append({"t": "title", "text": st[2:]})
            i += 1
            continue
        m = re.match(r"^%(author|affil|abstract|keywords)\s+(.*)$", st)
        if m:
            blocks.append({"t": m.group(1), "text": m.group(2)})
            i += 1
            continue
        if st == "%references":
            blocks.append({"t": "refshead"})
            in_refs = True
            i += 1
            continue
        m = re.match(r"^%table\s+(\w+)\s*\|\s*(.*?)\s*\|\s*([cf])$", st)
        if m:
            key, cap, w = m.groups()
            i += 1
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            assert len(rows) >= 3, (key, rows)
            blocks.append({"t": "table", "key": key, "caption": cap, "w": w, "rows": rows})
            continue
        m = re.match(r"^%eq\s+(\w+)\s*\|\s*(.*)$", st)
        if m:
            blocks.append({"t": "eq", "key": m.group(1), "text": m.group(2)})
            i += 1
            continue
        m = re.match(r"^(#{2,4})\s+(.*?)(?:\s+\{#(\w+)\})?$", st)
        if m:
            blocks.append({"t": "h", "lvl": len(m.group(1)) - 1, "text": m.group(2), "key": m.group(3)})
            in_refs = False
            i += 1
            continue
        m = re.match(r"^!\[fig:(\w+)\|(.*?)\|([cf]|full|col)\|(.*?)\]\((.*?)\)$", st)
        if m:
            key, cap, w, path, _ = m.groups()
            blocks.append({"t": "fig", "key": key, "caption": cap, "w": "f" if w in ("f", "full") else "c", "path": path})
            i += 1
            continue
        if st.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append({"t": "code", "text": "\n".join(code)})
            continue
        bm = re.match(r"^-\s+(.*)$", st)
        if bm:
            items = []
            while i < n and re.match(r"^-\s+", lines[i].strip()):
                items.append(re.sub(r"^-\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"t": "bullets", "items": items})
            continue
        nm = re.match(r"^(\d+)\.\s+(.*)$", st)
        if nm:
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"t": "numbers", "items": items})
            continue
        if in_refs and re.match(r"^\[\d+\]\s", st):
            blocks.append({"t": "ref", "text": st})
            i += 1
            continue
        para = [st]
        i += 1
        while (i < n and lines[i].strip() and not lines[i].strip().startswith(("#", "%", "![", "```", "|"))
               and not re.match(r"^(-|\d+\.)\s", lines[i].strip())):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"t": "para", "text": " ".join(para)})
    return blocks


blocks = parse(SRC.read_text(encoding="utf-8"))

# ---------------------------------------------------------------------------------------------- numbering
labels = {}   # key -> (full label, short label, bookmark)
l1 = l2 = l3 = 0
appendix_letter = {}
tab_n = fig_n = eq_n = 0
for b in blocks:
    if b["t"] == "h":
        text = b["text"]
        am = re.match(r"^Appendix ([A-Z]):\s*(.*)$", text)
        if am:
            b["appendix"] = am.group(1)
            b["title"] = am.group(2)
            labels["app" + am.group(1)] = ("Appendix " + am.group(1), am.group(1), "app_" + am.group(1))
            continue
        if b["lvl"] == 1:
            l1 += 1; l2 = l3 = 0
            b["num"] = ROMAN[l1] + "."
            lab = ROMAN[l1]
        elif b["lvl"] == 2:
            l2 += 1; l3 = 0
            b["num"] = chr(64 + l2) + "."
            lab = f"{ROMAN[l1]}-{chr(64 + l2)}"
        else:
            l3 += 1
            b["num"] = f"{l3})"
            lab = f"{ROMAN[l1]}-{chr(64 + l2)}{l3}"
        if b["key"]:
            labels[b["key"]] = ("Section " + lab, lab, "sec_" + b["key"])
    elif b["t"] == "table":
        tab_n += 1
        b["num"] = ROMAN[tab_n]
        labels["tab:" + b["key"]] = ("Table " + ROMAN[tab_n], ROMAN[tab_n], "tab_" + b["key"])
    elif b["t"] == "fig":
        fig_n += 1
        b["num"] = fig_n
        labels["fig:" + b["key"]] = (f"Fig. {fig_n}", str(fig_n), "fig_" + b["key"])
    elif b["t"] == "eq":
        eq_n += 1
        b["num"] = eq_n
        labels["eq:" + b["key"]] = (f"({eq_n})", str(eq_n), "eq_" + b["key"])
for b in blocks:
    if b["t"] == "ref":
        n_ = int(re.match(r"^\[(\d+)\]", b["text"]).group(1))
        labels[f"cite:{n_}"] = (f"[{n_}]", str(n_), f"ref_{n_}")

# ---------------------------------------------------------------------------------------------- document setup
doc = Document()
for st_name in ("Normal",):
    st = doc.styles[st_name]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0
for hn, size in (("Heading 1", 10), ("Heading 2", 10), ("Heading 3", 10), ("Title", 24)):
    s = doc.styles[hn]
    s.font.name = FONT
    s.font.size = Pt(size)
    s.font.bold = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
from docx.enum.style import WD_STYLE_TYPE
hl = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
hl.font.color.rgb = RGBColor(0x1F, 0x3F, 0xBF)
hl.font.underline = True
HYPER = hl.style_id

sec0 = doc.sections[0]
sec0.page_width, sec0.page_height = Cm(PAGE_W), Cm(PAGE_H)
sec0.left_margin = sec0.right_margin = Cm(MARGIN_LR)
sec0.top_margin, sec0.bottom_margin = Cm(1.9), Cm(2.2)


def set_cols(section, n):
    sp = section._sectPr
    for c in sp.findall(qn("w:cols")):
        sp.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    if n > 1:
        cols.set(qn("w:space"), str(int(GAP / 2.54 * 1440)))
        cols.set(qn("w:equalWidth"), "1")
    sp.append(cols)


def _fld(run, kind):
    el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)


def add_field(p, instr, size):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = p.add_run()
        r.font.size = Pt(size)
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        elif kind == "text":
            r.text = txt
        else:
            _fld(r, kind)


fp = sec0.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, "PAGE", 8)

set_cols(sec0, 1)
cur_cols = [1]
ncols_now = lambda: cur_cols[0]


def shrink_break_paragraph():
    """The paragraph that carries a section break should not take up a visible line."""
    p_elements = doc.element.body.findall(qn("w:p"))
    for pe in reversed(p_elements):
        if pe.find(qn("w:pPr")) is not None and pe.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            ppr = pe.find(qn("w:pPr"))
            sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0"); sp.set(qn("w:line"), "20"); sp.set(qn("w:lineRule"), "exact")
            ppr.insert(0, sp)
            rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "2"); rpr.append(sz)
            ppr.append(rpr)
            break


def want_cols(n):
    if cur_cols[0] == n:
        return
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(sec, n)
    cur_cols[0] = n
    shrink_break_paragraph()


_bm = [100]


def bookmark(p, name):
    _bm[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_bm[0])); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_bm[0]))
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(s)
    else:
        p._p.insert(0, s)
    p._p.append(e)


# ---------------------------------------------------------------------------------------------- inline text
TOK = re.compile(r"(\*\*.+?\*\*|`[^`]+?`|\*[^*\s][^*]*?\*|\{(?:sec|secn|fig|tab|app|cite|eq):[^}]+\}"
                 r"|https://github\.com/dorianbrown/rank_bm25|github\.com/Hanny658/FYP26-Proposals)")


def true_minus(t):
    return re.sub(r"(?:(?<=[\s\[(])|^)-(?=\d)", "−", t)


def make_hyperlink(p, text, anchor=None, url=None, size=None, bold=False, italic=False):
    h = OxmlElement("w:hyperlink")
    if anchor:
        h.set(qn("w:anchor"), anchor)
    else:
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), HYPER); rpr.append(rs)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def resolve(kind, key):
    if kind in ("sec", "secn"):
        lab = labels[key]
        return (lab[0] if kind == "sec" else lab[1]), lab[2]
    if kind == "fig":
        lab = labels["fig:" + key]; return lab[0], lab[2]
    if kind == "tab":
        lab = labels["tab:" + key]; return lab[0], lab[2]
    if kind == "app":
        lab = labels["app" + key]; return lab[0], lab[2]
    if kind == "cite":
        lab = labels["cite:" + key]; return lab[0], lab[2]
    if kind == "eq":
        lab = labels["eq:" + key]; return lab[0], lab[2]
    raise KeyError(kind)


def add_inline(p, text, size=None, bold=False, italic=False, mono=False, small_caps=False):
    for tok in TOK.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(p, tok[2:-2], size, True, italic, mono, small_caps)
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            r.font.name = MONO
            r.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            r.font.size = Pt((size or 10) - 1.5)
            r.bold = bold or None
            continue
        if tok.startswith("*") and tok.endswith("*") and len(tok) > 2 and not tok.startswith("**"):
            add_inline(p, tok[1:-1], size, bold, True, mono, small_caps)
            continue
        m = re.match(r"^\{(sec|secn|fig|tab|app|cite|eq):([^}]+)\}$", tok)
        if m:
            label, anchor = resolve(m.group(1), m.group(2))
            make_hyperlink(p, label, anchor=anchor, size=size, bold=bold, italic=italic)
            continue
        if tok.startswith("https://") or tok.startswith("github.com/"):
            make_hyperlink(p, tok, url=tok if tok.startswith("http") else "https://" + tok, size=size, bold=bold, italic=italic)
            continue
        r = p.add_run(true_minus(tok))
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        if small_caps:
            r.font.small_caps = True


def new_par(align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, indent=None, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if keep_next:
        pf.keep_with_next = True
    return p


# alignment choices taken from the edited Word file (the baseline): everything in a table is centred except a few
# long-text columns, body text is justified, and a few headings and equations carry justified alignment too
LEFT_COLS = {"t07": {0}, "t10": {3}, "t13": {2}, "t14": {5}}
JUSTIFIED_HEADINGS = {"prices", "fields", "own"}
JUSTIFIED_EQ = {"e_recall", "e_mrr", "e_ndcg"}
LEFT_PARA_PREFIX = "The command `python -m tests.wp4.eval_constraints --baseline`"


# ---------------------------------------------------------------------------------------------- tables
def split_row(line):
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def plain(c):
    c = re.sub(r"\{(?:sec|secn|fig|tab|app|cite|eq):[^}]+\}", lambda m: resolve(*m.group(0)[1:-1].split(":", 1))[0], c)
    return re.sub(r"[*`]", "", c)


def col_widths(header, rows, total):
    n = len(header)
    mins, weights = [], []
    for j in range(n):
        cells = [header[j]] + [r[j] for r in rows if j < len(r)]
        lens = [len(plain(c)) for c in cells]
        mx, avg = max(lens), sum(lens) / len(lens)
        weights.append(min(max(0.55 * mx + 0.45 * avg, 3.0), 46.0))
        longest = max(max((len(w) for w in plain(c).split()), default=0) for c in cells)
        mins.append(0.145 * longest + 0.25)
    if sum(mins) >= total:
        k = total / sum(mins)
        return [m * k for m in mins]
    spare = total - sum(mins)
    tw = sum(weights)
    return [m + spare * w / tw for m, w in zip(mins, weights)]


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_margins(table):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 15), ("left", 55), ("bottom", 15), ("right", 55)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblpr.append(mar)


def render_table(b):
    want_cols(1 if b["w"] == "f" else 2)
    total = FULL_W if b["w"] == "f" else COL_W
    cp = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, keep_next=True)
    add_inline(cp, f"TABLE {b['num']}", size=8, small_caps=True)
    bookmark(cp, "tab_" + b["key"])
    cp2 = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=0, after=3, keep_next=True)
    add_inline(cp2, b["caption"], size=8, small_caps=True)
    header = split_row(b["rows"][0])
    rows = [split_row(r) for r in b["rows"][2:]]
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    widths = col_widths(header, rows, total)
    fs = 7.5 if b["w"] == "c" else 8
    for c, h in enumerate(header):
        cell = t.rows[0].cells[c]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, h, size=fs, bold=True)
        shade(cell, "E7E7E7")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    ncol = len(header)
    col_max = [max(len(plain(r[j])) for r in rows if j < len(r)) for j in range(ncol)]
    centre = [j not in LEFT_COLS.get(b["key"], set()) for j in range(ncol)]
    for ri, row in enumerate(rows, start=1):
        trpr = t.rows[ri]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trpr.append(cs)
        for c, ctext in enumerate(row[:len(header)]):
            p = t.rows[ri].cells[c].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre[c] else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, ctext, size=fs)
    for row in t.rows:
        for c, cell in enumerate(row.cells):
            cell.width = Cm(widths[c])
    for c, col in enumerate(t.columns):
        col.width = Cm(widths[c])
    sp = new_par(before=0, after=4)
    sp.paragraph_format.line_spacing = Pt(4)


# ---------------------------------------------------------------------------------------------- block rendering
refs_started = False
for b in blocks:
    k = b["t"]
    if k == "title":
        want_cols(1)
        p = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=6, after=8)
        r = p.add_run(b["text"]); r.font.size = Pt(22)
    elif k == "author":
        p = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2)
        r = p.add_run(b["text"]); r.font.size = Pt(12)
    elif k == "affil":
        p = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0)
        add_inline(p, b["text"], size=9, italic=True)
    elif k == "abstract":
        want_cols(2)
        p = new_par(before=10, after=0)
        r = p.add_run("Abstract—"); r.bold = True; r.italic = True; r.font.size = Pt(9)
        add_inline(p, b["text"], size=9, bold=True)
    elif k == "keywords":
        p = new_par(before=4, after=0)
        r = p.add_run("Index Terms—"); r.bold = True; r.italic = True; r.font.size = Pt(9)
        add_inline(p, b["text"], size=9, bold=True)
    elif k == "h":
        want_cols(2)
        if b.get("appendix"):
            p = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep_next=True)
            p.style = doc.styles["Heading 1"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, f"Appendix {b['appendix']}: {b['title']}", size=10, small_caps=True)
            bookmark(p, "app_" + b["appendix"])
            continue
        if b["lvl"] == 1:
            p = doc.add_paragraph(style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12), Pt(6)
            add_inline(p, f"{b['num']} {b['text']}", size=10, small_caps=True)
        elif b["lvl"] == 2:
            p = doc.add_paragraph(style="Heading 2")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if b["key"] in JUSTIFIED_HEADINGS else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(3)
            add_inline(p, f"{b['num']} {b['text']}", size=10, italic=True)
        else:
            p = doc.add_paragraph(style="Heading 3")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(2)
            add_inline(p, f"{b['num']} {b['text']}", size=10, italic=True)
        p.paragraph_format.keep_with_next = True
        if b["key"]:
            bookmark(p, "sec_" + b["key"])
    elif k == "refshead":
        want_cols(2)
        p = doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12), Pt(6)
        add_inline(p, "References", size=10, small_caps=True)
        p.paragraph_format.keep_with_next = True
    elif k == "ref":
        want_cols(2)
        m = re.match(r"^\[(\d+)\]\s+(.*)$", b["text"])
        p = new_par(WD_ALIGN_PARAGRAPH.LEFT, before=0, after=3)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        add_inline(p, f"[{m.group(1)}]\t{m.group(2)}", size=8)
        bookmark(p, f"ref_{m.group(1)}")
    elif k == "para":
        want_cols(2)
        left = b["text"].startswith(LEFT_PARA_PREFIX)
        p = new_par(WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, indent=0.4)
        add_inline(p, b["text"])
    elif k in ("bullets", "numbers"):
        want_cols(2)
        for idx, it in enumerate(b["items"], start=1):
            p = new_par(before=0, after=2)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.6))
            add_inline(p, (f"{idx}.\t" if k == "numbers" else "•\t") + it)
    elif k == "eq":
        want_cols(2)
        p = new_par(WD_ALIGN_PARAGRAPH.JUSTIFY if b["key"] in JUSTIFIED_EQ else WD_ALIGN_PARAGRAPH.LEFT, before=3, after=5)
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Cm(COL_W / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Cm(COL_W), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        add_inline(p, b["text"], size=9)
        p.add_run(f"\t({b['num']})")
        bookmark(p, "eq_" + b["key"])
    elif k == "table":
        render_table(b)
    elif k == "fig":
        want_cols(1 if b["w"] == "f" else 2)
        width = {"pipeline": 17.0, "confirm": 17.0, "results": 12.0, "evalflow": 17.0}.get(b["key"], COL_W - 0.2)
        p = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, keep_next=True)
        p.add_run().add_picture(str((ROOT / b["path"]).resolve()), width=Cm(width))
        cp = new_par(WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
        add_inline(cp, f"Fig. {b['num']}. {b['caption']}", size=8)
        bookmark(cp, "fig_" + b["key"])
    elif k == "code":
        want_cols(1)
        p = new_par(WD_ALIGN_PARAGRAPH.LEFT, before=3, after=5)
        r = p.add_run(b["text"]); r.font.name = MONO; r.font.size = Pt(7.5)
        r.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)

doc.save(OUT)
print("saved", OUT, "| tables", tab_n, "| figures", fig_n, "| equations", eq_n, "| sections", l1)
